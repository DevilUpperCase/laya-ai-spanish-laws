from pathlib import Path
import csv
import hashlib
import re
from bs4 import BeautifulSoup
from openpyxl import load_workbook


def load_ine_snapshot(directory: Path):
    revision="INE-2026-01-01+captured-changes"
    entities={"ES":{"id":"ES","level":"state","parent_id":None,"name":"España","official_source_url":"https://www.ine.es/daco/daco42/codmun/cod_ccaa_provincia.htm","inventory_revision":revision,"verified":False}}
    soup=BeautifulSoup((directory/"cod_ccaa_provincia.htm").read_bytes(),"html.parser")
    for row in soup.select("table tr"):
        cells=[c.get_text(" ",strip=True) for c in row.find_all(["td","th"])]
        if len(cells)!=4 or not re.fullmatch(r"\d{2}",cells[0]):continue
        auto,name,province,prov_name=cells
        region="ES-AUTO-"+auto
        entities[region]={"id":region,"level":"autonomous_city" if auto in {"18","19"} else "autonomous_community",
             "parent_id":"ES","name":name,"official_source_url":"https://www.ine.es/daco/daco42/codmun/cod_ccaa_provincia.htm","inventory_revision":revision,"verified":False}
        if auto not in {"18","19"}:
            entities["ES-PRO-"+province]={"id":"ES-PRO-"+province,"level":"province","parent_id":region,"name":prov_name,
             "official_source_url":"https://www.ine.es/daco/daco42/codmun/cod_ccaa_provincia.htm","inventory_revision":revision,"verified":False}
    workbook=load_workbook(directory/"diccionario26.xlsx",read_only=True)
    rows=iter(workbook.active.values)
    title=str(next(rows)[0])
    if "2026" not in title:raise ValueError("Municipality workbook reference year mismatch")
    headers=next(rows)
    if tuple(headers[:5])!=("CODAUTO","CPRO","CMUN","DC","NOMBRE"):raise ValueError("Municipality schema changed")
    for auto,province,municipality,control,name in rows:
        auto,province,municipality=str(auto).zfill(2),str(province).zfill(2),str(municipality).zfill(3)
        identifier="ES-MUN-"+province+municipality
        if identifier in entities:raise ValueError("Duplicate municipality identity")
        parent="ES-AUTO-"+auto if auto in {"18","19"} else "ES-PRO-"+province
        if parent not in entities:raise ValueError("Municipality parent missing")
        entities[identifier]={"id":identifier,"level":"municipality","parent_id":parent,"name":name,
              "official_source_url":"https://www.ine.es/daco/daco42/codmun/diccionario26.xlsx","inventory_revision":revision,"verified":False}
    workbook.close()
    limitations=["Independent inventory review and foral institutional supplement pending."]
    changes=BeautifulSoup((directory/"codmun_anual.htm").read_bytes(),"html.parser")
    for table in changes.find_all("table"):
        table_rows=table.find_all("tr")
        heading=" ".join(r.get_text(" ",strip=True) for r in table_rows[:2])
        body=table_rows[2:]
        for row in body:
            cells=[c.get_text(" ",strip=True) for c in row.find_all("td")]
            if len(cells)<4 or not cells[0]:continue
            if "Denominación nueva" in heading or "Denominación corregida" in heading:
                match=re.fullmatch(r"(\d{2})\s+(\d{3})\s+\d",cells[2])
                if match:
                    identifier="ES-MUN-"+"".join(match.groups())
                    if identifier not in entities:raise ValueError("Change refers to absent municipality")
                    entities[identifier]["name"]=cells[1]
                    entities[identifier]["official_source_url"]="https://www.ine.es/daco/daco42/codmun/codmun_anual.htm"
            elif re.search(r"\d", " ".join(cells)):
                limitations.append("Municipality creation/deletion change requires manual reconciliation: "+cells[0])
    islands=load_workbook(directory/"26codislas.xlsx",read_only=True)
    for sheet in islands:
        rows=list(sheet.values)
        title=" ".join(str(c) for c in rows[0] if c)
        if "2026" not in title:
            limitations.append("Island workbook internal reference date differs from URL: "+title)
        header=next((index for index,row in enumerate(rows) if row[0]=="CPRO"),None)
        if header is None:raise ValueError("Island schema changed")
        for row in rows[header+1:]:
            province,island,name,municipality,control,mun_name=row[:6]
            if not province or island is None:continue
            identifier="ES-ISLAND-"+str(island)
            entities[identifier]={"id":identifier,"level":"island","parent_id":"ES-PRO-"+str(province).zfill(2),"name":str(name).strip(),
                 "official_source_url":"https://www.ine.es/daco/daco42/codmun/26codislas.xlsx","inventory_revision":title,"verified":False}
    islands.close()
    return list(entities.values()),list(dict.fromkeys(limitations))


def write_inventory(entities,path):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",encoding="utf-8",newline="") as file:
        writer=csv.DictWriter(file,fieldnames=["jurisdiction_id","level","parent_id","name","official_source_url"])
        writer.writeheader()
        for e in entities:writer.writerow({"jurisdiction_id":e["id"],**{k:e[k] for k in ("level","parent_id","name","official_source_url")}})
