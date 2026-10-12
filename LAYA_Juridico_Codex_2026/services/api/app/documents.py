import asyncio
from io import BytesIO
from pathlib import Path
import hashlib
import socket
import struct
from uuid import uuid4
from zipfile import ZipFile, BadZipFile
import pymupdf
from docx import Document
from defusedxml.ElementTree import fromstring
from .contracts import DocumentRecord, DocumentPage
from .errors import DomainError

DOCX="application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def validate_upload(data: bytes, mime: str):
    if mime == "application/pdf" and data.startswith(b"%PDF-"):
        return mime
    if mime != DOCX or not data.startswith(b"PK"):
        raise DomainError("UNSUPPORTED_DOCUMENT",415)
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos=archive.infolist()
            if len(infos)>5000 or sum(i.file_size for i in infos)>100*1024*1024:
                raise DomainError("ARCHIVE_LIMIT",413)
            if any(i.file_size>20*1024*1024 or i.file_size/max(i.compress_size,1)>100 for i in infos):
                raise DomainError("ARCHIVE_LIMIT",413)
            names={i.filename for i in infos}
            if not {"[Content_Types].xml","word/document.xml"}.issubset(names):
                raise DomainError("INVALID_DOCX",415)
            for info in infos:
                if ".." in Path(info.filename).parts or info.filename.startswith(("/","\\")) or "vbaproject" in info.filename.casefold():
                    raise DomainError("UNSAFE_DOCX",415)
                if info.filename.endswith(".rels"):
                    root=fromstring(archive.read(info))
                    if any(n.get("TargetMode")=="External" for n in root):
                        raise DomainError("EXTERNAL_DOCX_RELATIONSHIP",415)
        return mime
    except (BadZipFile,ValueError,KeyError):
        raise DomainError("INVALID_DOCX",415) from None


def scan_upload(data,host,port):
    if not host:
        return False  # Explicit dev limitation; production Settings requires a scanner.
    try:
        with socket.create_connection((host,port),timeout=10) as stream:
            stream.sendall(b"zINSTREAM\0")
            for offset in range(0,len(data),65536):
                chunk=data[offset:offset+65536]
                stream.sendall(struct.pack("!I",len(chunk))+chunk)
            stream.sendall(struct.pack("!I",0))
            result=stream.recv(4096)
            if b"FOUND" in result:
                raise DomainError("MALWARE_DETECTED",415)
            if not result.rstrip(b"\0\n").endswith(b"OK"):
                raise DomainError("MALWARE_SCANNER_UNAVAILABLE",503)
        return True
    except OSError:
        raise DomainError("MALWARE_SCANNER_UNAVAILABLE",503) from None


def extract_pages(data,mime,max_pages,max_chars,*,ocr=False,ocr_language="spa+cat+eus+glg"):
    pages=[]
    if mime == "application/pdf":
        with pymupdf.open(stream=data,filetype="pdf") as doc:
            if doc.is_encrypted:
                raise DomainError("ENCRYPTED_PDF",415)
            if len(doc)>max_pages or doc.xref_length()>100000:
                raise DomainError("DOCUMENT_LIMIT",413)
            for page in doc:
                text=page.get_text(sort=True)
                uncertain=False
                if not text.strip():
                    if not ocr:
                        raise DomainError("OCR_REQUIRED",409)
                    pixel_width,pixel_height=page.rect.width*200/72,page.rect.height*200/72
                    if pixel_width<=0 or pixel_height<=0 or max(pixel_width,pixel_height)>20000 or pixel_width*pixel_height>16_000_000:
                        raise DomainError("DOCUMENT_LIMIT",413)
                    tp=page.get_textpage_ocr(language=ocr_language,dpi=200,full=True)
                    text=page.get_text(textpage=tp)
                    uncertain=True
                pages.append({"page_number":len(pages)+1,"text":text,"ocr_uncertain":uncertain})
                if sum(len(p["text"]) for p in pages)>max_chars:
                    raise DomainError("DOCUMENT_LIMIT",413)
    else:
        doc=Document(BytesIO(data))
        # DOCX has no stable original page geometry: use explicit logical paragraphs/sections.
        paragraphs=[p.text for p in doc.paragraphs]
        paragraphs += [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        text="\n".join(paragraphs)
        if len(text)>max_chars:
            raise DomainError("DOCUMENT_LIMIT",413)
        pages=[{"page_number":1,"text":text,"ocr_uncertain":False}]
    if not pages or not any(p["text"].strip() for p in pages):
        raise DomainError("NO_EXTRACTABLE_TEXT",409)
    return pages


class DocumentService:
    def __init__(self,db,settings):
        self.db,self.settings=db,settings

    def object_path(self,key):
        root=self.settings.private_storage_path.resolve()
        path=(root/key).resolve()
        if not path.is_relative_to(root):
            raise DomainError("STORAGE_PATH_INVALID",503)
        return path

    async def upload(self,context,file,idempotency_key):
        if context.role == "reader":
            raise DomainError("WRITE_FORBIDDEN",403)
        data=await file.read(self.settings.upload_max_bytes+1)
        if len(data)>self.settings.upload_max_bytes:
            raise DomainError("UPLOAD_TOO_LARGE",413)
        if not data:
            raise DomainError("EMPTY_UPLOAD",415)
        mime=validate_upload(data,file.content_type)
        await asyncio.to_thread(scan_upload,data,self.settings.clamav_host,self.settings.clamav_port)
        filename=Path((file.filename or "document").replace("\\","/")).name[:200]
        if not filename or any(ord(ch)<32 for ch in filename):
            raise DomainError("FILENAME_INVALID",422)
        sha=hashlib.sha256(data).hexdigest()
        async with self.db.tenant_transaction(context) as conn:
            await conn.execute("SELECT pg_advisory_xact_lock(hashtextextended(%s,0))",(str(context.tenant_id)+":"+idempotency_key,))
            existing=await (await conn.execute("SELECT * FROM private.documents WHERE idempotency_key=%s",(idempotency_key,))).fetchone()
            if existing:
                if existing["deleted_at"] or existing["sha256"] != sha:
                    raise DomainError("IDEMPOTENCY_CONFLICT",409)
                return self._record(existing)
            document_id=uuid4()
            key=f"{context.tenant_id}/{document_id}/original"
            path=self.object_path(key)
            path.parent.mkdir(parents=True,exist_ok=True)
            try:
                await asyncio.to_thread(path.write_bytes,data)
                row=await (await conn.execute("""INSERT INTO private.documents(id,tenant_id,uploader,filename,mime_type,object_key,sha256,idempotency_key)
                 VALUES(%s,%s,%s,%s,%s,%s,%s,%s) RETURNING *""",(document_id,context.tenant_id,context.user_id,filename,mime,key,sha,idempotency_key))).fetchone()
                await conn.execute("INSERT INTO private.jobs(document_id,tenant_id,kind,deletion_epoch) VALUES(%s,%s,'extract',0)",(document_id,context.tenant_id))
                return self._record(row)
            except Exception:
                path.unlink(missing_ok=True)
                raise

    @staticmethod
    def _record(row,pages=None):
        return DocumentRecord(id=row["id"],tenant_id=row["tenant_id"],filename=row["filename"],mime_type=row["mime_type"],
              parse_status=row["parse_status"],created_at=row["created_at"],failure_code=row["failure_code"],pages=pages or [],folder_id=row.get("folder_id"),trashed_at=row.get("trashed_at"))

    async def get(self,context,document_id,*,trash=False):
        async with self.db.tenant_transaction(context) as conn:
            row=await (await conn.execute("SELECT * FROM private.documents WHERE id=%s AND deleted_at IS NULL AND (trashed_at IS NOT NULL)=%s",(document_id,trash))).fetchone()
            if row is None:
                raise DomainError("DOCUMENT_NOT_FOUND",404)
            pages=await (await conn.execute("SELECT page_number,text,ocr_uncertain FROM private.chunks WHERE document_id=%s ORDER BY page_number",(document_id,))).fetchall()
            return self._record(row,[DocumentPage(**p) for p in pages])

    async def delete(self,context,document_id,*,trash_only=False):
        if context.role == "reader":
            raise DomainError("WRITE_FORBIDDEN",403)
        async with self.db.tenant_transaction(context) as conn:
            row=await (await conn.execute("SELECT * FROM private.documents WHERE id=%s FOR UPDATE",(document_id,))).fetchone()
            if row is None:
                raise DomainError("DOCUMENT_NOT_FOUND",404)
            if row["deleted_at"]:
                return
            if trash_only and not row["trashed_at"]:return
            epoch=row["deletion_epoch"]+1
            from .deletion_journal import record_intent
            # Write before committing the tombstone: a crash may conservatively reapply
            # the authorized deletion, but cannot resurrect an acknowledged deletion.
            await asyncio.to_thread(record_intent,self.settings.deletion_journal_path,context.tenant_id,document_id,epoch)
            await conn.execute("UPDATE private.documents SET deleted_at=now(),purge_due_at=now()+interval '24 hours',deletion_epoch=%s,parse_status='deleted' WHERE id=%s",(epoch,document_id))
            await conn.execute("UPDATE private.jobs SET status='cancelled',lease_token=NULL WHERE document_id=%s AND kind IN ('extract','index')",(document_id,))
            await conn.execute("INSERT INTO private.deletion_records VALUES(%s,%s,now(),now()+interval '24 hours',NULL,now()+interval '30 days',%s)",(document_id,context.tenant_id,epoch))
            await conn.execute("INSERT INTO private.jobs(document_id,tenant_id,kind,deletion_epoch) VALUES(%s,%s,'purge',%s)",(document_id,context.tenant_id,epoch))

    async def download(self,context,document_id,*,trash=False):
        async with self.db.tenant_transaction(context) as conn:
            row=await (await conn.execute("SELECT object_key,mime_type,filename FROM private.documents WHERE id=%s AND deleted_at IS NULL AND (trashed_at IS NOT NULL)=%s",(document_id,trash))).fetchone()
            if row is None:
                raise DomainError("DOCUMENT_NOT_FOUND",404)
            return row

    async def stream(self,context,document_id,key,*,trash=False):
        path=self.object_path(key)
        with path.open("rb") as file:
            while True:
                async with self.db.tenant_transaction(context) as conn:
                    row=await (await conn.execute("SELECT id FROM private.documents WHERE id=%s AND deleted_at IS NULL AND (trashed_at IS NOT NULL)=%s FOR SHARE",(document_id,trash))).fetchone()
                    if row is None:
                        return
                    block=file.read(65536)
                    if not block:
                        return
                    yield block
