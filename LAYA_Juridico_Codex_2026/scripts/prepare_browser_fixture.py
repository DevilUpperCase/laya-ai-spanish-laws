"""Clearly synthetic private PDF for browser workflow checks; never ingest as legislation."""
from pathlib import Path
import pymupdf
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
if __name__=="__main__":
    path=ROOT/"private-data/browser-synthetic.pdf"
    path.parent.mkdir(exist_ok=True)
    document=pymupdf.open()
    document.new_page().insert_text((50,50),"Synthetic private test only: vivienda y arrendamiento. Not legal evidence.")
    document.save(path)
    word_path=ROOT/"private-data/browser-synthetic.docx"
    word=Document()
    word.add_paragraph("Synthetic private Unicode test only: vivienda y habitación.")
    word.add_paragraph("Una marca 🏠 precede al arrendamiento doméstico. No es evidencia legal.")
    word.save(word_path)
    print("Synthetic private browser fixture prepared.")
